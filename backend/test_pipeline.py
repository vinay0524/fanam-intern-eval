"""
Unit test script to verify DocxParser, TextChunker, and SparseBM25Service
without needing live external API keys.
"""
import io
from docx import Document
from services.docx_parser import DocxParser
from services.text_chunker import TextChunker
from services.sparse_service import SparseBM25Service
from config import settings

def run_local_tests():
    print("--- 1. Testing Settings & Validation ---")
    status = settings.validate_keys()
    print(f"Key status detected: configured={status['configured']}, missing={status['missing_keys']}")

    print("\n--- 2. Testing DOCX Creation & Parsing ---")
    # Create in-memory docx
    doc = Document()
    doc.add_heading("Q3 Financial Performance Report", 0)
    doc.add_paragraph("Total revenue for the quarter reached $12.4 billion, representing a 14% year-over-year increase.")
    doc.add_paragraph("Operating margins expanded to 28.5% driven by cloud infrastructure efficiencies.")
    
    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Segment"
    hdr_cells[1].text = "Revenue ($B)"
    hdr_cells[2].text = "YoY Growth"
    row_cells = table.rows[1].cells
    row_cells[0].text = "Cloud Services"
    row_cells[1].text = "5.8"
    row_cells[2].text = "+26%"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    extracted = DocxParser.parse_docx_bytes(file_bytes, filename="sample_q3.docx")
    print(f"Extracted Text Preview:\n{extracted}")
    assert "12.4 billion" in extracted
    assert "Cloud Services" in extracted
    print("DocxParser passed successfully!")

    print("\n--- 3. Testing Text Chunker ---")
    chunker = TextChunker(chunk_size=120, chunk_overlap=30)
    chunks = chunker.chunk_text(extracted, source_filename="sample_q3.docx")
    print(f"Generated {len(chunks)} chunks:")
    for c in chunks:
        print(f"  [{c['chunk_id']}] (len={c['char_count']}): {c['text'][:60]}...")
    assert len(chunks) > 0
    print("TextChunker passed successfully!")

    print("\n--- 4. Testing Sparse BM25 Retrieval Vector Generator ---")
    sparse_encoder = SparseBM25Service()
    chunk_texts = [c["text"] for c in chunks]
    sparse_vectors = sparse_encoder.encode_documents(chunk_texts)
    
    assert len(sparse_vectors) == len(chunks)
    for i, sv in enumerate(sparse_vectors):
        print(f"  Chunk {i} sparse vector: {len(sv['indices'])} active terms")
        assert len(sv["indices"]) == len(sv["values"])
        assert all(idx >= 0 for idx in sv["indices"])

    # Test query sparse vector
    query_sparse = sparse_encoder.encode_query("What was the cloud revenue growth?")
    print(f"Query sparse vector: indices={query_sparse['indices']}, values={query_sparse['values']}")
    assert len(query_sparse["indices"]) > 0
    print("SparseBM25Service passed successfully!")

    print("\nALL OFFLINE PIPELINE TESTS PASSED!")

if __name__ == "__main__":
    run_local_tests()

