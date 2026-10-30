import io
import csv
from typing import Tuple, List, Dict, Any
from docx import Document
from pypdf import PdfReader
import openpyxl

class DocumentParser:
    """
    Unified multi-format file loader supporting:
    - .docx (Microsoft Word)
    - .pdf (PDF Documents)
    - .csv (Spreadsheet data)
    - .txt / .md (Plain text files)
    - .xlsx (Excel Workbooks)
    """

    @classmethod
    def parse_file(cls, file_bytes: bytes, filename: str) -> str:
        ext = filename.lower().split(".")[-1]

        if ext == "docx":
            return cls.parse_docx(file_bytes)
        elif ext == "pdf":
            return cls.parse_pdf(file_bytes)
        elif ext == "csv":
            return cls.parse_csv(file_bytes)
        elif ext in ["txt", "md", "log"]:
            return cls.parse_text(file_bytes)
        elif ext in ["xlsx", "xls"]:
            return cls.parse_excel(file_bytes)
        else:
            raise ValueError(f"Unsupported file format: .{ext}. Supported formats: .docx, .pdf, .csv, .txt, .xlsx")

    @staticmethod
    def parse_docx(file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        parts = []

        # Paragraphs & headings
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # Tables (e.g. financial metrics, balance sheets)
        for table_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                dedup_cells = []
                for cell in row_cells:
                    if not dedup_cells or cell != dedup_cells[-1]:
                        dedup_cells.append(cell)
                if any(dedup_cells):
                    table_rows.append(" | ".join(dedup_cells))
            
            if table_rows:
                parts.append(f"\n[Table {table_idx + 1}]\n" + "\n".join(table_rows) + "\n")

        return "\n\n".join(parts)

    @staticmethod
    def parse_pdf(file_bytes: bytes) -> str:
        reader = PdfReader(io.BytesIO(file_bytes))
        parts = []

        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                parts.append(f"[Page {page_idx + 1}]\n{page_text.strip()}")

        return "\n\n".join(parts)

    @staticmethod
    def parse_csv(file_bytes: bytes) -> str:
        # Try UTF-8 first, fallback to latin-1
        try:
            content = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            content = file_bytes.decode("latin-1")

        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if not rows:
            return ""

        headers = rows[0]
        parts = [f"Columns: {', '.join(headers)}"]

        # Format rows into readable tabular text
        table_lines = [" | ".join(headers)]
        table_lines.append("-" * len(table_lines[0]))

        for row in rows[1:]:
            if any(cell.strip() for cell in row):
                table_lines.append(" | ".join(row))

        parts.append("\n".join(table_lines))
        return "\n\n".join(parts)

    @staticmethod
    def parse_text(file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8").strip()
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1").strip()

    @staticmethod
    def parse_excel(file_bytes: bytes) -> str:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        parts = []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sheet_rows = []
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty rows
                cells = [str(c).strip() if c is not None else "" for c in row]
                if any(cells):
                    sheet_rows.append(" | ".join(cells))
            
            if sheet_rows:
                parts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(sheet_rows))

        return "\n\n".join(parts)

